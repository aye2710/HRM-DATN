import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

SRC_MD = r"c:\Users\truclh\Desktop\HRM-ĐATN\docs\BAO_CAO_DO_AN_TOT_NGHIEP_HRM.md"
OUT_DOCX_1 = r"c:\Users\truclh\Desktop\HRM-ĐATN\docs\Ho_So_DATN_Word\BAO_CAO_DO_AN_TOT_NGHIEP_HRM.docx"
OUT_DOCX_2 = r"c:\Users\truclh\Desktop\HRM-ĐATN\docs\Ho_So_DATN_Word\DATN-14-Thuyet-minh.docx"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def set_run_font(run, font_name="Times New Roman", font_size=Pt(13)):
    run.font.name = font_name
    run.font.size = font_size
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}"'
        f' w:cs="{font_name}" w:eastAsia="{font_name}"/>'
    )
    rPr.append(rFonts)

def clean_xml_string(s):
    if not s:
        return ""
    # Chuyển đổi toàn bộ thẻ HTML br sang dấu ngắt dòng thực tế trong Word
    s = re.sub(r'<br\s*/?>', '\n', s, flags=re.IGNORECASE)
    s = re.sub(r'&nbsp;', ' ', s, flags=re.IGNORECASE)
    s = re.sub(r'&amp;', '&', s, flags=re.IGNORECASE)
    s = re.sub(r'&lt;', '<', s, flags=re.IGNORECASE)
    s = re.sub(r'&gt;', '>', s, flags=re.IGNORECASE)
    # Chuyển đổi toàn bộ ký hiệu LaTeX sang Unicode
    s = re.sub(r'\$\\rightarrow\$', '→', s)
    s = re.sub(r'\\rightarrow', '→', s)
    s = re.sub(r'\$\\leftarrow\$', '←', s)
    s = re.sub(r'\\leftarrow', '←', s)
    s = re.sub(r'\$\\ge\$', '≥', s)
    s = re.sub(r'\\ge\b', '≥', s)
    s = re.sub(r'\$\\le\$', '≤', s)
    s = re.sub(r'\\le\b', '≤', s)
    s = re.sub(r'\$\\times\$', '×', s)
    s = re.sub(r'\\times\b', '×', s)
    # Loại bỏ ký tự điều khiển không hợp lệ trong XML
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

def add_formatted_runs(paragraph, text, base_font="Times New Roman", base_size=Pt(13), base_color=RGBColor(0,0,0)):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.color.rgb = base_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = clean_xml_string(token[2:-2])
            set_run_font(run, "Times New Roman", base_size)
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = clean_xml_string(token[1:-1])
            set_run_font(run, "Times New Roman", base_size)
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = clean_xml_string(token[1:-1])
            # Đồng bộ toàn bộ chữ về Times New Roman (trừ sơ đồ)
            set_run_font(run, "Times New Roman", base_size)
            run.font.color.rgb = RGBColor(140, 20, 20)
            run.bold = True
        else:
            run.text = clean_xml_string(token)
            set_run_font(run, "Times New Roman", base_size)

def convert():
    print(f"Reading markdown: {SRC_MD}")
    with open(SRC_MD, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = docx.enum.section.WD_ORIENT.PORTRAIT
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(20)
    sec.left_margin = Mm(30)
    sec.right_margin = Mm(20)

    # Đặt Times New Roman cho tất cả các Styles cơ bản trong document
    for style_name in ['Normal', 'List Bullet', 'List Number']:
        if style_name in doc.styles:
            st = doc.styles[style_name]
            st.font.name = 'Times New Roman'
            st.font.size = Pt(13)
            rPr = st._element.get_or_add_rPr()
            rFonts = parse_xml(
                f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"'
                f' w:cs="Times New Roman" w:eastAsia="Times New Roman"/>'
            )
            rPr.append(rFonts)

    style = doc.styles['Normal']
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.space_before = Pt(0)

    in_code_block = False
    code_block_lines = []
    code_lang = ""

    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        parsed_rows = []
        for l in tbl_lines:
            l = l.strip()
            if not l.startswith('|'):
                continue
            cols = [c.strip() for c in l.split('|')[1:-1]]
            # check if it's separator row
            if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                continue
            parsed_rows.append(cols)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        if num_cols == 0:
            return

        table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        col_widths = [Inches(6.3 / num_cols)] * num_cols

        for row_idx, row_data in enumerate(parsed_rows):
            row = table.rows[row_idx]
            set_row_cant_split(row)
            if row_idx == 0:
                set_repeat_header(row)
            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell_text = clean_xml_string(cell_text)
                
                # Format cell text
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                
                if row_idx == 0:
                    set_cell_background(cell, "F0F4F8")
                    run = p.add_run(cell_text)
                    set_run_font(run, "Times New Roman", Pt(11.5))
                    run.font.color.rgb = RGBColor(10, 30, 80)
                    run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "FAFAFA")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    add_formatted_runs(p, cell_text, base_font="Times New Roman", base_size=Pt(11))
                    if re.match(r'^\d+(\.\d+)?%?$', cell_text) or cell_text in ['Có', 'Không', 'Đạt', 'Bắt buộc', 'Tùy chọn', 'Pass', 'Fail']:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add tiny spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.space_before = Pt(0)

    def flush_code_block(c_lines, lang):
        if not c_lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.3)
        set_cell_background(cell, "F6F8FA")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        # border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'<w:left w:val="single" w:sz="16" w:space="0" w:color="0969DA"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        # Heading indicating code/diagram type (giữ Consolas riêng cho sơ đồ/mã nguồn)
        header_text = f"[{lang.upper() if lang else 'MÃ NGUỒN / SƠ ĐỒ THỰC THI'}]"
        hrun = p.add_run(header_text + "\n")
        set_run_font(hrun, "Consolas", Pt(9.5))
        hrun.font.color.rgb = RGBColor(9, 105, 218)
        hrun.bold = True

        for i, code_l in enumerate(c_lines):
            crun = p.add_run(code_l + ("\n" if i < len(c_lines)-1 else ""))
            set_run_font(crun, "Consolas", Pt(9.5))
            crun.font.color.rgb = RGBColor(36, 41, 47)
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    idx = 0
    total_lines = len(lines)
    while idx < total_lines:
        raw_line = clean_xml_string(lines[idx])
        stripped = raw_line.strip()

        # Handle Code block
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block(code_block_lines, code_lang)
                in_code_block = False
                code_block_lines = []
                code_lang = ""
            else:
                if in_table:
                    flush_table(table_lines)
                    in_table = False
                    table_lines = []
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_block_lines = []
            idx += 1
            continue

        if in_code_block:
            code_block_lines.append(raw_line.rstrip('\r\n'))
            idx += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            idx += 1
            continue
        elif in_table:
            flush_table(table_lines)
            in_table = False
            table_lines = []

        # Empty lines
        if not stripped:
            idx += 1
            continue

        # Headings (Times New Roman)
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            # If major chapter, page break
            if any(h_text.startswith(prefix) for prefix in ["CHƯƠNG ", "Chương ", "LỜI CẢM ƠN", "MỤC LỤC", "KẾT LUẬN"]):
                doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(h_text)
            set_run_font(run, "Times New Roman", Pt(18))
            run.font.color.rgb = RGBColor(0, 32, 96)
            run.bold = True
            idx += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(h_text)
            set_run_font(run, "Times New Roman", Pt(15))
            run.font.color.rgb = RGBColor(0, 51, 153)
            run.bold = True
            idx += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h_text)
            set_run_font(run, "Times New Roman", Pt(13.5))
            run.font.color.rgb = RGBColor(20, 20, 20)
            run.bold = True
            idx += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h_text)
            set_run_font(run, "Times New Roman", Pt(13))
            run.font.color.rgb = RGBColor(40, 40, 40)
            run.bold = True
            run.italic = True
            idx += 1
            continue

        if stripped.startswith('##### '):
            h_text = stripped[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h_text)
            set_run_font(run, "Times New Roman", Pt(12.5))
            run.font.color.rgb = RGBColor(60, 60, 60)
            run.bold = True
            idx += 1
            continue

        # Bullet list (Times New Roman)
        if stripped.startswith('- ') or stripped.startswith('* '):
            content = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(13))
            idx += 1
            continue

        # Numbered list (Times New Roman)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            content = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run_num = p.add_run(f"{num}. ")
            set_run_font(run_num, "Times New Roman", Pt(13))
            run_num.bold = True
            add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(13))
            idx += 1
            continue

        # Blockquote (Times New Roman)
        if stripped.startswith('> '):
            content = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, content, base_font="Times New Roman", base_size=Pt(13), base_color=RGBColor(80, 80, 80))
            for run in p.runs:
                run.italic = True
            idx += 1
            continue

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        # Regular paragraph (Times New Roman)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_runs(p, stripped, base_font="Times New Roman", base_size=Pt(13))
        idx += 1

    # Cleanup any trailing table/code block
    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_block_lines, code_lang)

    print(f"Saving copy to Word docx: {OUT_DOCX_2}")
    try:
        doc.save(OUT_DOCX_2)
        print(f"Saved successfully: {OUT_DOCX_2}")
    except Exception as e:
        print(f"Could not save {OUT_DOCX_2}: {e}")

    print(f"Saving to Word docx: {OUT_DOCX_1}")
    try:
        doc.save(OUT_DOCX_1)
        print(f"Saved successfully: {OUT_DOCX_1}")
    except PermissionError:
        alt_path = OUT_DOCX_1.replace(".docx", "_TimesNewRoman.docx")
        print(f"File {OUT_DOCX_1} is currently open in Word! Saved to {alt_path} instead.")
        doc.save(alt_path)
    print("Done converting thesis to docx!")

if __name__ == "__main__":
    convert()
